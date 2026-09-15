"""
Export settled bets into the Excel tracker.

The tracker already computes yield, win rate, payoff ratio, Kelly, Sharpe,
t-stat and p-value from columns A-R of "1. Bet Entry", with the formulas
pre-filled down to row 295. So this writes ONLY the input columns and
leaves every formula alone:

    A  Entry Date        <- placed_at
    E  SPORT             <- sport
    F  Bet Type          <- market
    G  Wager Amount      <- stake
    H  Odds (decimal)    <- price actually taken
    I  Win (1/0)         <- settlement
    M  Pick              <- e.g. "Patrick Mahomes Over 249.5"
    N  Bet Specifics     <- the matchup
    O  Strategy          <- book and the model's EV at bet time

B, C, J, K, P, Q and R are formulas and are never touched.

Only won and lost bets are exported. The sheet's P&L formula is
`=IF(I=1, G*(H-1), -G)`, which treats a blank Win cell as a loss -- so
writing a pending bet would book a phantom loss. Pushes, voids and
half-settled bets have no representation in that formula at all; they are
skipped and counted, rather than fudged into a number that would corrupt
the t-stat.
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

TEMPLATE = Path(__file__).parent / "templates" / "tracker_template.xlsx"
SHEET = "1. Bet Entry"
FIRST_ROW = 2

# 1-indexed columns, matching the tracker's own layout.
COL_DATE, COL_SPORT, COL_TYPE = 1, 5, 6
COL_STAKE, COL_ODDS, COL_WIN = 7, 8, 9
COL_PICK, COL_SPECIFICS, COL_STRATEGY = 13, 14, 15

EXPORTABLE = {"won": 1, "lost": 0}

SPORT_LABEL = {
    "basketball_nba": "NBA",
    "americanfootball_nfl": "NFL",
    "icehockey_nhl": "NHL",
    "baseball_mlb": "MLB",
    "soccer_epl": "EPL",
    "basketball_ncaab": "NCAAB",
    "americanfootball_ncaaf": "NCAAF",
    "mma_mixed_martial_arts": "MMA",
}


class TrackerExportError(RuntimeError):
    pass


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise TrackerExportError(
            "Excel export needs openpyxl. Run: pip install openpyxl"
        ) from exc
    return __import__("openpyxl")


def describe_pick(bet) -> str:
    line = "" if bet["line"] is None else f" {bet['line']:g}"
    return f"{bet['selection'] or ''} {bet['side'] or ''}{line}".strip()


def describe_strategy(bet) -> str:
    book = (bet["book"] or "").title()
    ev = bet["ev_at_bet"]
    if ev is None:
        return f"{book} (manual)"
    return f"{book} {ev:+.1%} EV"


def _parse_date(value) -> datetime | None:
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    # Excel cannot store a timezone-aware datetime.
    return dt.replace(tzinfo=None)


def export_tracker(
    bets: Sequence[Any],
    output_path: str | Path,
    template_path: str | Path | None = None,
) -> dict[str, Any]:
    """
    Write settled bets into a copy of the tracker.

    Returns a summary dict: how many rows were written, how many were
    skipped and why.
    """
    openpyxl = _require_openpyxl()
    template = Path(template_path or TEMPLATE)
    if not template.exists():
        raise TrackerExportError(f"tracker template missing at {template}")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, output_path)

    wb = openpyxl.load_workbook(output_path)
    if SHEET not in wb.sheetnames:
        raise TrackerExportError(
            f"'{SHEET}' not found in the template; sheets are {wb.sheetnames}"
        )
    ws = wb[SHEET]

    exportable, skipped = [], {}
    for bet in bets:
        status = (bet["status"] or "").lower()
        if status in EXPORTABLE:
            exportable.append(bet)
        else:
            skipped[status] = skipped.get(status, 0) + 1

    capacity = ws.max_row - FIRST_ROW + 1
    if len(exportable) > capacity:
        _extend_formula_rows(ws, FIRST_ROW, ws.max_row, len(exportable))

    for i, bet in enumerate(exportable):
        row = FIRST_ROW + i
        placed = _parse_date(bet["placed_at"])
        if placed is not None:
            ws.cell(row, COL_DATE).value = placed.date()
            ws.cell(row, COL_DATE).number_format = "yyyy-mm-dd"
        ws.cell(row, COL_SPORT).value = SPORT_LABEL.get(bet["sport"], bet["sport"])
        ws.cell(row, COL_TYPE).value = _market_label(bet["market"])
        ws.cell(row, COL_STAKE).value = bet["stake"]
        ws.cell(row, COL_ODDS).value = bet["price"]
        ws.cell(row, COL_WIN).value = EXPORTABLE[(bet["status"] or "").lower()]
        ws.cell(row, COL_PICK).value = describe_pick(bet)
        ws.cell(row, COL_SPECIFICS).value = bet["matchup"]
        ws.cell(row, COL_STRATEGY).value = describe_strategy(bet)

    wb.save(output_path)

    return {
        "path": output_path,
        "written": len(exportable),
        "skipped": skipped,
        "total_staked": sum(b["stake"] for b in exportable),
        "pnl": sum(
            b["stake"] * (b["price"] - 1) if b["status"] == "won" else -b["stake"]
            for b in exportable
        ),
    }


def _market_label(key: str | None) -> str:
    from .report import pretty_market

    return pretty_market(key) if key else ""


def _extend_formula_rows(ws, first_row: int, last_row: int, needed: int) -> None:
    """
    Copy the formula columns down so the sheet can hold more bets than the
    template was built for. Only the computed columns are copied; the input
    columns are left empty for the caller to fill.
    """
    from copy import copy

    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter

    formula_cols = [2, 3, 10, 11, 16, 17, 18]  # B C J K P Q R
    target_last = first_row + needed - 1
    for row in range(last_row + 1, target_last + 1):
        for col in formula_cols:
            src = ws.cell(last_row, col)
            dst = ws.cell(row, col)
            if isinstance(src.value, str) and src.value.startswith("="):
                # Translator shifts every relative reference by the row
                # delta. A naive string replace would leave the running
                # total's back-reference (=J295+R294) pointing at the wrong
                # row, silently breaking the rolling P&L column.
                letter = get_column_letter(col)
                dst.value = Translator(
                    src.value, origin=f"{letter}{last_row}"
                ).translate_formula(f"{letter}{row}")
            if src.has_style:
                dst._style = copy(src._style)
